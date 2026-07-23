from __future__ import annotations

import json
import sqlite3
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFD", str(value or "").strip().upper())
    return "".join(char for char in text if unicodedata.category(char) != "Mn")


def _text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(microsecond=0).isoformat(sep=" ")
    text = str(value).strip()
    return text or None


def _get(row: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        if key in row:
            return row[key]
    return None


def _last_movement(row: dict[str, Any]) -> str | None:
    date_value = _text(_get(row, "FECHA DE ULTIMO MOVIMIENTO", "FECHA ULTIMO MOVIMIENTO"))
    time_value = _text(_get(row, "HORA DE ULTIMO MOVIMIENTO", "HORA ULTIMO MOVIMIENTO"))
    if not date_value:
        return None
    return f"{date_value} {time_value}" if time_value else date_value


def import_orders_xlsx(database_path: Path, xlsx_path: Path) -> dict[str, int]:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [_normalize_header(cell) for cell in next(rows, ())]
    records = []
    for values in rows:
        row = {headers[index]: values[index] if index < len(values) else None for index in range(len(headers))}
        order_id = _text(_get(row, "ID", "ORDER ID"))
        if not order_id:
            continue
        records.append((order_id, row))

    with sqlite3.connect(database_path) as connection:
        connection.execute("PRAGMA foreign_keys=ON")
        run = connection.execute("INSERT INTO import_runs(source_path, rows_count) VALUES(?, ?)", (str(xlsx_path), len(records)))
        run_id = run.lastrowid
        created = 0
        updated = 0
        for order_id, row in records:
            guide = _text(_get(row, "NUMERO GUIA", "GUIA"))
            status = _text(_get(row, "ESTATUS", "STATUS"))
            carrier = _text(_get(row, "TRANSPORTADORA", "CARRIER"))
            movement = _last_movement(row)
            payload = json.dumps({key: _text(value) for key, value in row.items()}, ensure_ascii=False, sort_keys=True)
            exists = connection.execute("SELECT 1 FROM orders WHERE order_id=?", (order_id,)).fetchone() is not None
            connection.execute(
                """INSERT INTO orders(order_id, guide, status, carrier, last_movement_at, raw_json)
                   VALUES(?, ?, ?, ?, ?, ?)
                   ON CONFLICT(order_id) DO UPDATE SET guide=excluded.guide, status=excluded.status,
                   carrier=excluded.carrier, last_movement_at=excluded.last_movement_at,
                   raw_json=excluded.raw_json, updated_at=CURRENT_TIMESTAMP""",
                (order_id, guide, status, carrier, movement, payload),
            )
            connection.execute(
                """INSERT INTO order_snapshots(import_run_id, order_id, status, guide, carrier, last_movement_at, raw_json)
                   VALUES(?, ?, ?, ?, ?, ?, ?)""",
                (run_id, order_id, status, guide, carrier, movement, payload),
            )
            created += 0 if exists else 1
            updated += 1 if exists else 0
    return {"rows": len(records), "created": created, "updated": updated}
